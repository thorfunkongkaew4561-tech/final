from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PROJECT_NAME_TH = (
    'การพัฒนาระบบปัญญาประดิษฐ์และระบบสนับสนุนการตัดสินใจสำหรับการบริหาร'
    'การปล่อยเช่าคอนโดมิเนียมในรูปแบบเว็บแอปพลิเคชัน'
)
PROJECT_NAME_EN = (
    'Development of an Artificial Intelligence and Decision Support System '
    'for Condominium Rental Management as a Web Application'
)
OUTPUT_PATH = (
    Path(__file__).resolve().parents[2]
    / 'deliverables'
    / 'Condo_Rental_AI_Project_Report.docx'
)

doc = Document()

sections = doc.sections
for section in sections:
    section.top_margin = Inches(1.5)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)

style_normal = doc.styles['Normal']
font_normal = style_normal.font
font_normal.name = 'TH SarabunPSK'
font_normal.size = Pt(16)
font_normal.color.rgb = RGBColor(0, 0, 0) 
style_normal.paragraph_format.first_line_indent = Inches(0.6) 
style_normal.paragraph_format.line_spacing = 1.15
style_normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'TH SarabunPSK')
style_normal.element.rPr.rFonts.set(qn('w:cs'), 'TH SarabunPSK')

style_h1 = doc.styles['Heading 1']
font_h1 = style_h1.font
font_h1.name = 'TH SarabunPSK'
font_h1.size = Pt(18)
font_h1.bold = True
font_h1.color.rgb = RGBColor(0, 0, 0)
style_h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_h1.paragraph_format.space_after = Pt(12)
style_h1.element.rPr.rFonts.set(qn('w:cs'), 'TH SarabunPSK')

style_h2 = doc.styles['Heading 2']
font_h2 = style_h2.font
font_h2.name = 'TH SarabunPSK'
font_h2.size = Pt(16)
font_h2.bold = True
font_h2.color.rgb = RGBColor(0, 0, 0)
style_h2.paragraph_format.space_before = Pt(12)
style_h2.paragraph_format.space_after = Pt(6)
style_h2.element.rPr.rFonts.set(qn('w:cs'), 'TH SarabunPSK')

# ฟังก์ชันสำหรับสร้างข้อความจัดกลางแบบไม่มี Indent
def add_centered_paragraph(text, size=18, bold=True, space_after=0, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    run = p.add_run(text)
    run.font.name = 'TH SarabunPSK'
    run.font.size = Pt(size)
    run.font.bold = bold
    return p


def format_result_table(table):
    """จัดรูปแบบตารางผลลัพธ์ให้สอดคล้องกับรายงาน."""

    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()

            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '1F4E78' if row_index == 0 else 'FFFFFF')
            properties.append(shading)

            borders = properties.first_child_found_in('w:tcBorders')
            if borders is None:
                borders = OxmlElement('w:tcBorders')
                properties.append(borders)
            for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                border = OxmlElement(f'w:{edge}')
                border.set(qn('w:val'), 'single')
                border.set(qn('w:sz'), '4')
                border.set(qn('w:color'), 'D9D9D9')
                borders.append(border)

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = Inches(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = 'TH SarabunPSK'
                    run.font.size = Pt(14)
                    run.font.color.rgb = RGBColor(
                        255 if row_index == 0 else 0,
                        255 if row_index == 0 else 0,
                        255 if row_index == 0 else 0,
                    )
                    run.font.bold = row_index == 0

# ==========================================
# 1. หน้าปกนอก (Cover Page)
# ==========================================
add_centered_paragraph('[ ใส่ตราสัญลักษณ์มหาวิทยาลัยศิลปากร ตรงนี้ ]', size=16, bold=False, space_after=20, space_before=20)
add_centered_paragraph('รายงานการปฏิบัติงานสหกิจศึกษา', size=20, space_after=40)

add_centered_paragraph(PROJECT_NAME_TH, size=18)
add_centered_paragraph(f'({PROJECT_NAME_EN})', size=16, space_after=40)

add_centered_paragraph('นางสาวอาทิตยา มีจันทร์   รหัสนักศึกษา [เติมรหัส]', size=18)
add_centered_paragraph('นางสาวทอฝัน [นามสกุล]   รหัสนักศึกษา [เติมรหัส]', size=18, space_after=60)

add_centered_paragraph('รายงานฉบับนี้เป็นส่วนหนึ่งของรายวิชา [รหัสวิชา] สหกิจศึกษา', size=18)
add_centered_paragraph('สาขาวิชาวิทยาการข้อมูล คณะวิทยาศาสตร์', size=18)
add_centered_paragraph('มหาวิทยาลัยศิลปากร', size=18)
add_centered_paragraph('ภาคการศึกษา [ระบุภาค] ปีการศึกษา [ระบุปี]', size=18)

doc.add_page_break()

# ==========================================
# 2. หน้าปกใน (Title Page)
# ==========================================
add_centered_paragraph('[ ใส่ตราสัญลักษณ์มหาวิทยาลัยศิลปากร ตรงนี้ ]', size=16, bold=False, space_after=20, space_before=20)
add_centered_paragraph('รายงานการปฏิบัติงานสหกิจศึกษา', size=20, space_after=40)

add_centered_paragraph(PROJECT_NAME_TH, size=18)
add_centered_paragraph(f'({PROJECT_NAME_EN})', size=16, space_after=40)

add_centered_paragraph('นางสาวอาทิตยา มีจันทร์   รหัสนักศึกษา [เติมรหัส]', size=18)
add_centered_paragraph('นางสาวทอฝัน [นามสกุล]   รหัสนักศึกษา [เติมรหัส]', size=18, space_after=60)

add_centered_paragraph('ปฏิบัติงาน ณ', size=18)
add_centered_paragraph('[ ชื่อบริษัท / สถานประกอบการ ]', size=18)
add_centered_paragraph('[ ที่อยู่บริษัท ]', size=18)

doc.add_page_break()

# ==========================================
# เนื้อหา (เหมือนเดิม)
# ==========================================

doc.add_heading('กิตติกรรมประกาศ\n(Acknowledgement)', level=1)
doc.add_paragraph('[ ใส่ข้อความแสดงความขอบคุณผู้ที่ให้ความช่วยเหลือ และร่วมมือในการปฏิบัติงาน... ]')
doc.add_page_break()

doc.add_heading('บทคัดย่อ\n(Abstract)', level=1)
doc.add_paragraph('การปฏิบัติงานสหกิจศึกษาครั้งนี้เป็นส่วนหนึ่งของโครงงานพัฒนาระบบปัญญาประดิษฐ์และระบบสนับสนุนการตัดสินใจสำหรับการบริหารการปล่อยเช่าคอนโดมิเนียม โดยใช้ข้อมูลจำลองของโครงการ Kave ขอบเขตงานในรายงานฉบับนี้เน้นการพัฒนาโมเดลแนะนำราคาค่าเช่า (Rental Price Recommendation) และ Price Engine ส่วนโมเดลพยากรณ์สภาวะตลาดและอัตราการเข้าพักล่วงหน้า (Market & Occupancy Forecasting) เป็นความรับผิดชอบของเพื่อนร่วมโครงงาน และจะนำผลมาเชื่อมต่อกับ What-if Simulation และ Web Application เมื่อรวมระบบ')
doc.add_paragraph('ในส่วนของแบบจำลองราคาค่าเช่า ผู้จัดทำได้เปรียบเทียบอัลกอริทึม Stepwise Multiple Linear Regression, Gradient Boosting Regressor, Random Forest Regressor, K-Nearest Neighbors Regressor และ Artificial Neural Network โดยใช้ Rolling Time Validation 4 รอบภายในชุดฝึกสอน เพื่อลดความเสี่ยงจากการนำข้อมูลอนาคตมาใช้ในการพัฒนาโมเดล ผลการศึกษาพบว่า Random Forest มี Validation MAE ต่ำที่สุด 713.92 บาท และ Test MAE ต่ำที่สุด 1,020.73 บาท จึงถูกเลือกเป็นโมเดลหลักของ Price Engine ทั้งนี้ผลลัพธ์เป็นค่าประมาณและควรใช้ร่วมกับข้อมูลสภาพห้องและตลาด')
doc.add_page_break()

doc.add_heading('สารบัญ', level=1)
doc.add_paragraph('[ กดอัปเดตสารบัญอัตโนมัติใน Word ]')
doc.add_page_break()

# บทที่ 1
doc.add_heading('บทที่ 1\nบทนำ', level=1)
doc.add_heading('1.1 บทนำ หรือที่มาและความสำคัญ', level=2)
doc.add_paragraph('ธุรกิจการปล่อยเช่าอสังหาริมทรัพย์ โดยเฉพาะคอนโดมิเนียมที่ตั้งอยู่ใกล้กับสถาบันการศึกษา (เช่น โครงการแบรนด์ Kave) มีความท้าทายอย่างมากในเรื่องของการตั้งราคาค่าเช่า หากตั้งราคาสูงเกินไปก็อาจทำให้เสียโอกาสในการปล่อยเช่า เกิดปัญหาห้องว่างเป็นระยะเวลานาน (High Vacancy Rate) แต่ถ้าตั้งราคาต่ำเกินไปก็ทำให้ผู้ลงทุนสูญเสียรายได้ที่ควรจะได้ ปัจจุบันการตัดสินใจตั้งราคาหรือทำโปรโมชั่นมักจะใช้เพียงประสบการณ์ของผู้ดูแลโครงการ ทำให้ไม่สามารถปรับตัวตามความต้องการของตลาดที่เปลี่ยนไปในแต่ละช่วงเวลาได้อย่างทันท่วงที')
doc.add_paragraph('จากปัญหาดังกล่าว ผู้จัดทำจึงเห็นถึงความสำคัญของการนำข้อมูล (Data) ทั้งในอดีตและปัจจุบัน มาวิเคราะห์ร่วมกับเทคนิค Machine Learning เพื่อสร้างระบบที่สามารถประเมินและแนะนำราคาค่าเช่าที่เหมาะสมได้แบบอัตโนมัติ โดยวิเคราะห์จากลักษณะเฉพาะของห้องพักแต่ละห้อง (เช่น ขนาดพื้นที่ ชั้น ทิศทางวิว) และยังสามารถจำลองสถานการณ์ล่วงหน้า (What-if Simulation) เพื่อให้ผู้บริหารมองเห็นภาพรวมของรายได้และตัดสินใจได้อย่างมีประสิทธิภาพมากขึ้น')

doc.add_heading('1.2 วัตถุประสงค์', level=2)
p_obj = doc.add_paragraph()
p_obj.paragraph_format.first_line_indent = Inches(0)
p_obj.paragraph_format.left_indent = Inches(0.6)
p_obj.add_run('1. เพื่อพัฒนาโมเดลแนะนำราคาค่าเช่าที่เหมาะสม (Rental Price Recommendation) โดยพิจารณาจากปัจจัยแวดล้อมและสภาวะตลาด\n')
p_obj.add_run('2. เพื่อพัฒนาโมเดลพยากรณ์สภาวะตลาดและอัตราการเข้าพักล่วงหน้า (Market & Occupancy Forecasting)\n')
p_obj.add_run('3. เพื่อพัฒนาระบบ What-if Simulation สำหรับจำลองสถานการณ์การปรับราคาและประเมินผลกระทบต่อรายได้\n')
p_obj.add_run('4. เพื่อพัฒนาระบบสนับสนุนการตัดสินใจด้วย AI ผ่าน Web Application')

doc.add_heading('1.3 ขอบเขตของการทำงาน', level=2)
doc.add_paragraph('ในโครงงานนี้ใช้ข้อมูลจำลองอ้างอิงจากโครงการ Kave โดยโมเดลราคาเช่าแบ่งข้อมูลตามลำดับเวลา (Chronological Split) กำหนดให้ชุดฝึกสอน (Train) อยู่ในช่วงปี 2023-2025 ชุดตรวจสอบ (Validation) อยู่ในช่วงเดือนมกราคมถึงมีนาคม 2026 และชุดทดสอบ (Test) อยู่ในเดือนเมษายนถึงมิถุนายน 2026 เพื่อให้กระบวนการประเมินสอดคล้องกับการนำโมเดลไปใช้กับข้อมูลอนาคต และลดความเสี่ยงจาก Data Leakage ขอบเขตงานของอาทิตยาครอบคลุมการสกัดฟีเจอร์ราคาเช่า การฝึกและจูนโมเดล Price Engine การเชื่อมต่อ What-if และ Web Application')

doc.add_heading('1.4 ขั้นตอนการดำเนินงาน', level=2)
p_step = doc.add_paragraph()
p_step.paragraph_format.first_line_indent = Inches(0)
p_step.paragraph_format.left_indent = Inches(0.6)
p_step.add_run('1. ตรวจข้อมูลทั้ง 5 ตาราง รวมถึงลำดับเวลาและนิยามอัตราการเข้าพักใน weekly_market\n')
p_step.add_run('2. เตรียม features สำหรับราคาเช่าและ features ย้อนหลังสำหรับการพยากรณ์ตลาดและอัตราการเข้าพัก\n')
p_step.add_run('3. ฝึกและประเมินโมเดลราคาเช่า ส่วนโมเดลตลาดและอัตราการเข้าพักต้องทดลองและประเมินแยก target และ horizon\n')
p_step.add_run('4. เชื่อม Price Engine กับผลพยากรณ์รายสัปดาห์ใน What-if เพื่อคำนวณรายได้ระดับโครงการ\n')
p_step.add_run('5. พัฒนาหน้าแนะนำราคา หน้าพยากรณ์ หน้าจำลองสถานการณ์ และหน้าสรุปผลบน Streamlit')

doc.add_heading('1.5 ประโยชน์ที่ได้รับจากการปฏิบัติงาน', level=2)
doc.add_paragraph('การปฏิบัติงานครั้งนี้ทำให้ผู้จัดทำได้เรียนรู้กระบวนการพัฒนา Data Product จริงๆ ตั้งแต่ขั้นตอนการจัดการข้อมูล การเลือกใช้โมเดลทางสถิติให้เหมาะสมกับบริบทของปัญหา (เช่น การใช้ Rolling Time Folds แทนการ Random Split ปกติ) และการออกแบบซอฟต์แวร์ที่นำผลลัพธ์จาก Machine Learning ไปแสดงผลให้ผู้ใช้ทั่วไปเข้าใจได้ง่าย')
doc.add_page_break()

# บทที่ 2
doc.add_heading('บทที่ 2\nทฤษฎีที่เกี่ยวข้อง', level=1)
doc.add_paragraph('การพัฒนาโมเดลแนะนำราคาค่าเช่าอสังหาริมทรัพย์ อาศัยแนวคิดและทฤษฎีพื้นฐานทางสถิติดังต่อไปนี้')

doc.add_heading('2.1 ปัจจัยที่มีอิทธิพลต่อราคาอสังหาริมทรัพย์', level=2)
doc.add_paragraph('ในการประเมินมูลค่าค่าเช่า ปัจจัยที่เกี่ยวข้องมักแบ่งออกเป็น 2 กลุ่มหลัก ได้แก่ ปัจจัยทางกายภาพตัวห้อง (เช่น ขนาดพื้นที่ใช้สอย ชั้น การตกแต่งเฟอร์นิเจอร์ วิวห้อง) และปัจจัยด้านทำเลที่ตั้ง (เช่น ระยะห่างจากมหาวิทยาลัย จำนวนสิ่งอำนวยความสะดวกในโครงการ) ซึ่งสอดคล้องกับงานวิจัยที่ศึกษาอัตราค่าเช่าหอพักในพื้นที่รอบมหาวิทยาลัยธรรมศาสตร์ ศูนย์รังสิต ที่พบว่าระยะทางและส่วนกลางมีผลอย่างมีนัยสำคัญต่ออัตราค่าเช่า')

doc.add_heading('2.2 การพยากรณ์ราคาด้วย Multiple Linear Regression', level=2)
doc.add_paragraph('สมการถดถอยเชิงเส้นพหุคูณ เป็นวิธีการพื้นฐานที่ใช้หาความสัมพันธ์ระหว่างตัวแปรอิสระหลายตัวกับตัวแปรตาม ในบริบทนี้คือการหาว่าฟีเจอร์ต่างๆ ส่งผลบวกหรือลบต่อราคาเช่ากี่บาท ซึ่งมีข้อดีคือสามารถอธิบายผลลัพธ์ได้อย่างชัดเจน (Interpretability) โดยเทคนิค Stepwise/Forward Selection จะช่วยคัดกรองเฉพาะตัวแปรที่ส่งผลกระทบจริงๆ เข้ามาในสมการ')

doc.add_heading('2.3 การพยากรณ์ด้วย Gradient Boosting Regressor', level=2)
doc.add_paragraph('แม้วิธี Linear Regression จะอธิบายผลได้ดี แต่อาจไม่สามารถจับความสัมพันธ์แบบไม่เป็นเส้นตรง (Non-linear) ระหว่างตัวแปรได้ Gradient Boosting จึงถูกนำมาใช้เป็นหนึ่งในโมเดลเปรียบเทียบ โดยอ้างอิงแนวทางการประยุกต์ Machine Learning เพื่อพยากรณ์ราคาอสังหาริมทรัพย์ของ Hernes, Tutak และ Siewiera (2024)')

doc.add_heading('2.4 การพยากรณ์ด้วย Random Forest Regressor', level=2)
doc.add_paragraph('Random Forest เป็นอัลกอริทึมแบบ Ensemble ที่สร้าง Decision Tree หลายต้นจากตัวอย่างข้อมูลและชุดตัวแปรที่แตกต่างกัน แล้วนำผลการพยากรณ์ของทุกต้นมาเฉลี่ย วิธีนี้ช่วยจับความสัมพันธ์แบบไม่เป็นเส้นตรงและลดความแปรปรวนของต้นไม้เดี่ยว ในโครงงานนี้มีการปรับค่าจำนวนต้นไม้ ความลึกของต้นไม้ จำนวนข้อมูลขั้นต่ำในการแบ่งโหนด จำนวนข้อมูลขั้นต่ำในใบไม้ และสัดส่วนตัวแปรที่สุ่มใช้ในแต่ละจุดแบ่ง')

doc.add_heading('2.5 การพยากรณ์ด้วย K-Nearest Neighbors Regressor', level=2)
doc.add_paragraph('K-Nearest Neighbors Regressor พยากรณ์ราคาจากตัวอย่างที่มีลักษณะใกล้เคียงกับห้องเป้าหมาย โดยต้องปรับสเกลตัวแปรเชิงตัวเลขและเข้ารหัสตัวแปรเชิงกลุ่มก่อนคำนวณระยะห่าง โครงงานนี้จูนจำนวนเพื่อนบ้าน วิธีให้น้ำหนัก และรูปแบบการคำนวณระยะห่างด้วย Rolling Time Validation เช่นเดียวกับโมเดลอื่น')

doc.add_heading('2.6 การพยากรณ์ด้วย Artificial Neural Network', level=2)
doc.add_paragraph('Artificial Neural Network ในโครงงานนี้ใช้ Multi-layer Perceptron Regressor เพื่อเรียนรู้ความสัมพันธ์แบบไม่เป็นเส้นตรงระหว่างคุณลักษณะห้อง โครงการ ช่วงเวลา และราคาเช่า มีการจูนขนาดชั้นซ่อน ค่า regularization และ learning rate โดยใช้ข้อมูล Train เท่านั้น')
doc.add_page_break()

# บทที่ 3
doc.add_heading('บทที่ 3\nรายละเอียดของงานที่ปฏิบัติ', level=1)
doc.add_paragraph('ในบทนี้จะอธิบายถึงรายละเอียดการปฏิบัติงานในส่วนของโมเดลแนะนำราคาค่าเช่า (Rental Price) เป็นหลัก ตั้งแต่การเตรียมข้อมูลจนถึงการเขียนโปรแกรมนำโมเดลไปใช้งาน')

doc.add_heading('3.1 การสกัดตัวแปร (Feature Engineering)', level=2)
doc.add_paragraph('ในขั้นตอนนี้ ผู้จัดทำได้ทำการรวมข้อมูลจาก 3 ตารางหลัก ได้แก่ ตารางข้อมูลห้องพัก (units), ข้อมูลโครงการ (projects) และข้อมูลการเช่า (leases) เข้าด้วยกัน จากนั้นดึงฟีเจอร์ที่สำคัญออกมาทั้งหมด 13 ตัวแปร ได้แก่ ชั้น (floor), ขนาดพื้นที่ (size_sqm), ประเภทห้อง (room_type), ทิศทางวิว (view), การตกแต่ง (furnished), ระยะเวลาสัญญา (lease_months), สถานะการต่อสัญญา (is_renewal), ชื่อโครงการ (project_name), มหาวิทยาลัยใกล้เคียง (university), ระยะห่างจากมหาลัย (distance_to_campus_m), จำนวนสิ่งอำนวยความสะดวก (facility_count), ปีและเดือนที่เริ่มเช่า (start_year, start_month) เพื่อใช้เป็นอินพุต (X) สำหรับการทำนายราคา (y)')

doc.add_heading('3.2 การฝึกสอนและคัดเลือกโมเดล (Model Training and Selection)', level=2)
doc.add_paragraph('การฝึกสอนโมเดลไม่ได้ใช้วิธีสุ่มแบ่งข้อมูลแบบปกติ เนื่องจากข้อมูลมีลำดับเวลา ผู้จัดทำจึงแบ่งข้อมูลเป็น Train, Validation และ Test ตามวันที่เริ่มสัญญา จากนั้นใช้ Rolling Time Validation 4 รอบภายในชุด Train เพื่อเลือกฟีเจอร์และปรับ Hyperparameters โดยไม่ใช้ Validation หลักหรือ Test ในขั้นตอนดังกล่าว โมเดลที่เปรียบเทียบคือ Stepwise Multiple Linear Regression, Gradient Boosting Regressor, Random Forest Regressor, K-Nearest Neighbors Regressor และ Artificial Neural Network')
doc.add_paragraph('กำหนดให้ Validation MAE เป็นเกณฑ์หลักในการเลือกโมเดลก่อนเปิดดู Test ผลการคัดเลือกพบว่า Random Forest มี Validation MAE ต่ำที่สุด จึงถูกบันทึกเป็นโมเดลหลัก โดยค่าที่เลือกคือ n_estimators เท่ากับ 400, max_depth เท่ากับ 20, min_samples_split เท่ากับ 5, min_samples_leaf เท่ากับ 4 และ max_features เท่ากับ 0.7')

doc.add_heading('3.3 การสร้าง Price Engine', level=2)
doc.add_paragraph('หลังจากได้โมเดลผู้ชนะ ผู้จัดทำได้เขียนโมดูล Price Engine ซึ่งมีหน้าที่รับข้อมูลคุณสมบัติของห้องพักจากผู้ใช้งาน นำไปประมวลผลผ่านโมเดล แล้วส่งค่าราคาแนะนำออกมา (Recommended Rent) โดยระบบจะไม่เพียงแค่ให้ราคาตัวเลขโดดๆ แต่จะประมวลผลช่วงราคาต่ำสุด (Lower Rent) และราคาสูงสุด (Upper Rent) โดยอ้างอิงจาก Error Margin ของโมเดล ทำให้ผู้ใช้งานมีความยืดหยุ่นในการนำราคาไปใช้ตัดสินใจจริง')

doc.add_heading('3.4 การออกแบบ Market & Occupancy Forecasting', level=2)
doc.add_paragraph('วัตถุประสงค์ข้อ 2 ใช้ข้อมูล weekly_market ระดับโครงการรายสัปดาห์ พยากรณ์จำนวนผู้ค้นหา จำนวนประกาศเปิด ค่ามัธยฐานราคาเสนอเช่า และอัตราการเข้าพักล่วงหน้า 1, 4 และ 12 สัปดาห์ ต้องตรวจสอบนิยาม occupancy และจำนวนห้องพร้อมให้เช่าก่อนนำไปคำนวณรายได้')
doc.add_paragraph('การเตรียมข้อมูลใช้ lag และค่าเฉลี่ยย้อนหลังแยกโครงการ โดยทราบข้อมูลถึง forecast origin เท่านั้น แบ่ง Train Validation และ Test ตาม target date ของแต่ละ horizon ใช้ rolling-origin validation และเปรียบเทียบกับ baseline ค่าล่าสุดหรือฤดูกาลเดียวกัน รายงาน MAE และ RMSE แยกตัวแปรและช่วงพยากรณ์ สำหรับ occupancy แสดงความคลาดเคลื่อนเป็นจุดเปอร์เซ็นต์')
doc.add_paragraph('แบบจำลอง Market & Occupancy Forecasting ตามวัตถุประสงค์ข้อ 2 เป็นความรับผิดชอบของเพื่อนร่วมโครงงาน จึงไม่รวมโค้ดฝึก จูน หรือประเมินโมเดลดังกล่าวไว้ในส่วนงานนี้ เมื่อรวมระบบจะรับไฟล์โมเดลและ metadata ที่ระบุฟีเจอร์ ตัวแปรเป้าหมาย ช่วงพยากรณ์ และนิยามอัตราการเข้าพัก เพื่อนำผลตลาดและเส้นทางอัตราการเข้าพักรายสัปดาห์ไปใช้ใน What-if Simulation เพื่อคำนวณรายได้ระดับโครงการ โดยไม่แปลงเป็นระยะเวลาปล่อยเช่ารายห้อง')
doc.add_page_break()

doc.add_heading('3.5 การปรับ What-if และ Web Application', level=2)
doc.add_paragraph('ระบบ What-if รับเส้นทาง occupancy ที่พยากรณ์รายสัปดาห์ของโครงการเดียวกัน คำนวณรายได้จากราคาเช่ารายเดือนคูณจำนวนห้องพร้อมให้เช่า คูณ occupancy และสัดส่วนเดือนของแต่ละสัปดาห์ แล้วหักส่วนลดรวมในกรอบเวลานั้น การเปลี่ยนราคาใช้ occupancy เดิมเป็นสมมติฐาน จนกว่าจะพัฒนาแบบจำลองการตอบสนองต่อราคาเพิ่มเติม ผลคำนวณนี้เป็นรายได้ประมาณการ ไม่ใช่กำไรหรือผลเชิงเหตุและผล')
doc.add_paragraph('Web Application แบ่งเป็นส่วนแนะนำราคา พยากรณ์ตลาดและอัตราการเข้าพัก What-if และสรุปผล โดยแสดงโครงการ forecast origin และช่วงพยากรณ์ให้ชัดเจน โมเดลราคาเดิมยังไม่รับค่าตลาดล่วงหน้าเป็น feature จนกว่าจะฝึกและประเมินใหม่')
doc.add_page_break()

# บทที่ 4
doc.add_heading('บทที่ 4\nผลการดำเนินงาน', level=1)

doc.add_heading('4.1 ประสิทธิภาพแบบจำลอง Rental Price', level=2)
doc.add_paragraph('ผลการเปรียบเทียบแสดงว่า Random Forest มี Validation MAE ต่ำที่สุด 713.92 บาท และ Test MAE ต่ำที่สุด 1,020.73 บาท จึงได้รับเลือกตามเกณฑ์หลักที่กำหนดไว้ก่อนเปิดชุด Test อย่างไรก็ตาม Gradient Boosting มี Test RMSE ต่ำกว่าและ Test R² สูงกว่าเล็กน้อย จึงควรพิจารณาหลายตัวชี้วัดร่วมกัน และไม่สรุปว่า Random Forest ดีกว่าในทุกมิติ')

result_table = doc.add_table(rows=1, cols=5)
result_table.autofit = True
headers = ['โมเดล', 'ชุดข้อมูล', 'MAE', 'RMSE', 'R²']
for index, text in enumerate(headers):
    result_table.rows[0].cells[index].text = text

result_rows = [
    ['Stepwise MLR', 'Validation', '770.84', '1,051.46', '0.7186'],
    ['Gradient Boosting', 'Validation', '759.60', '1,021.25', '0.7345'],
    ['Random Forest', 'Validation', '713.92', '955.94', '0.7674'],
    ['KNN', 'Validation', '808.07', '1,111.78', '0.6854'],
    ['ANN', 'Validation', '870.11', '1,143.05', '0.6674'],
    ['Stepwise MLR', 'Test', '1,113.72', '1,395.80', '0.5854'],
    ['Gradient Boosting', 'Test', '1,032.37', '1,314.12', '0.6325'],
    ['Random Forest', 'Test', '1,020.73', '1,330.66', '0.6232'],
    ['KNN', 'Test', '1,117.42', '1,439.12', '0.5592'],
    ['ANN', 'Test', '1,069.96', '1,356.02', '0.6086'],
]
for values in result_rows:
    cells = result_table.add_row().cells
    for index, text in enumerate(values):
        cells[index].text = text
format_result_table(result_table)

doc.add_paragraph('เมื่อวิเคราะห์ Feature Importance ของ Random Forest พบว่ากลุ่มประเภทห้อง โดยเฉพาะ studio มีความสำคัญสูงสุด รองลงมาคือขนาดห้องและระยะห่างจากมหาวิทยาลัย ค่า Feature Importance แสดงว่าตัวแปรมีส่วนช่วยในการแบ่งข้อมูลของต้นไม้มากเพียงใด แต่ไม่สามารถใช้ยืนยันความสัมพันธ์เชิงเหตุและผลได้')

doc.add_heading('4.2 ผลลัพธ์จากระบบแนะนำราคา (Price Engine)', level=2)
doc.add_paragraph('Price Engine ส่งคืนราคาเช่าที่โมเดลคาดการณ์และช่วงความคลาดเคลื่อนอ้างอิงจากค่า P75 Absolute Error บน Validation ซึ่งเท่ากับ 990.23 บาท เมื่อตรวจสอบกับ Test พบว่า Margin ดังกล่าวครอบคลุมค่าจริง 58.80% ต่ำกว่าเป้าหมาย 75% จึงไม่ควรเรียกผลนี้ว่า Prediction Interval หรือช่วงความเชื่อมั่น ระบบแสดงข้อความกำกับว่าราคาเป็นค่าประมาณและควรพิจารณาสภาพห้องและตลาดเพิ่มเติม')

doc.add_heading('4.3 สถานะการพัฒนา Market & Occupancy Forecasting', level=2)
doc.add_paragraph('ส่วนพยากรณ์ตลาดและอัตราการเข้าพักอยู่ในขั้นกำหนดขอบเขต ยังไม่มีผลฝึกหรือคะแนนประเมินของโมเดลใหม่ จึงต้องทำ EDA เตรียม features ทดลองโมเดล และประเมินแยก horizon ก่อนเชื่อมระบบ')
doc.add_paragraph('ผลการประเมิน Market & Occupancy Forecasting จะนำมาจากส่วนงานของเพื่อนร่วมโครงงานเมื่อการพัฒนาและการตรวจสอบโมเดลเสร็จสมบูรณ์ รายงานฉบับนี้จึงไม่รายงานค่าประสิทธิภาพแทนหรือประมาณผลของโมเดลดังกล่าว')
doc.add_page_break()

# บทที่ 5
doc.add_heading('บทที่ 5\nสรุปปัญหา และข้อเสนอแนะ', level=1)

doc.add_heading('5.1 สรุปผลการปฏิบัติงาน', level=2)
doc.add_paragraph('ส่วนงาน Rental Price Recommendation ของโครงงานได้นำข้อมูลจำลองของธุรกิจเช่าคอนโดมาพัฒนาและเปรียบเทียบแบบจำลองราคาค่าเช่า 5 วิธี ผลการคัดเลือกด้วย Validation MAE พบว่า Random Forest เหมาะสมที่สุดสำหรับใช้เป็นโมเดลหลักของ Price Engine ในข้อมูลชุดนี้ โดยระบบนำผลพยากรณ์ไปแสดงเป็นราคาเช่าโดยประมาณพร้อมข้อความกำกับข้อจำกัด ทั้งนี้โมเดลทำนายค่าเช่าตามสัญญา โดยใช้วันที่เริ่มสัญญาเป็นตัวแทนช่วงเวลา')

doc.add_heading('5.2 ปัญหาและอุปสรรค', level=2)
doc.add_paragraph('เนื่องจากข้อมูลที่ได้รับเป็นเพียงข้อมูลจำลองและตัวแปรบางส่วนยังขาดความครบถ้วน (เช่น ข้อมูลด้านเศรษฐกิจ อัตราดอกเบี้ย ฤดูกาลเปิดปิดภาคเรียน) การฝึกสอนโมเดลจึงยังมีข้อจำกัดที่ทำให้โมเดลไม่สามารถตอบสนองต่อสภาวะเศรษฐกิจจริงได้อย่างสมบูรณ์ 100%')

doc.add_heading('5.3 ข้อเสนอแนะ', level=2)
doc.add_paragraph('ควรมีการเชื่อมโยงฐานข้อมูลการปล่อยเช่าจริงแบบ Real-time เข้ามาแทนที่ข้อมูลจำลอง และควรศึกษาการทำ A/B Testing ในการใช้งานจริง เพื่อตรวจสอบว่าเมื่อผู้ดูแลระบบนำราคาที่โมเดลแนะนำไปใช้ จะสามารถลดระยะเวลาห้องว่างได้ดีกว่าการตั้งราคาแบบเดิมมากน้อยเพียงใด')
doc.add_page_break()

# บรรณานุกรม
doc.add_heading('บรรณานุกรม', level=1)
p_bib = doc.add_paragraph()
p_bib.paragraph_format.first_line_indent = Inches(0)
p_bib.paragraph_format.left_indent = Inches(0.6)
p_bib.add_run('1. [Paper ในโฟลเดอร์] การพยากรณ์อัตราค่าเช่าที่อยู่อาศัยประเภทหอพักหรืออพาร์ทเม้นท์ในพื้นที่โดยรอบสถาบันการศึกษา กรณีศึกษา มหาวิทยาลัยธรรมศาสตร์ ศูนย์รังสิต\n')
p_bib.add_run('2. Hernes, M., Tutak, M., & Siewiera, A. (2024). Prediction of residential real estate price on primary market using machine learning. Procedia Computer Science, 246, 3142-3147.\n')
p_bib.add_run('3. A Machine Learning Approach Using Spatially Explicit K-Nearest Neighbors for House Price Predictions.\n')
p_bib.add_run('4. Appropriate Rental Price Prediction for Condominiums in Pattaya, Thailand Using Artificial Neural Network Approach.\n')
p_bib.add_run('5. Rental Price Prediction of Boarding Houses in Batam City Using Linear Regression and Random Forest Algorithms.')
doc.add_page_break()

# ภาคผนวก
doc.add_heading('ภาคผนวก', level=1)
doc.add_paragraph('ภาพหน้าจอ Web Application (Streamlit)')

OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT_PATH)
print(f'Created: {OUTPUT_PATH}')
